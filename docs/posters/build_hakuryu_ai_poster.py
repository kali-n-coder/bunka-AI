from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "hakuryu_ai_visitor_poster.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D7DCE8", size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=130, start=160, bottom=130, end=160):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=None, bold=False, color=None):
    run.font.name = "Noto Sans JP"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans JP")
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, size=12, bold=False, color="102033"):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run


def set_paragraph(paragraph, align=None, before=0, after=6, line=1.05):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_card(doc, title, body, fill="FFFFFF"):
    table = doc.add_table(rows=1, cols=1)
    table.allow_autofit = False
    table.columns[0].width = Cm(18.2)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    set_paragraph(p, before=0, after=2)
    add_text(p, title, size=14, bold=True, color="0F2742")
    p2 = cell.add_paragraph()
    set_paragraph(p2, before=0, after=0, line=1.12)
    add_text(p2, body, size=10.5, color="26384D")
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    style = doc.styles["Normal"]
    style.font.name = "Noto Sans JP"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans JP")
    style.font.size = Pt(10.5)

    # Header band
    header = doc.add_table(rows=1, cols=1)
    header.allow_autofit = False
    header.columns[0].width = Cm(18.2)
    hcell = header.cell(0, 0)
    set_cell_shading(hcell, "E8ECF5")
    set_cell_border(hcell, color="E8ECF5", size="4")
    set_cell_margins(hcell, top=260, bottom=260, start=260, end=260)
    p = hcell.paragraphs[0]
    set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_text(p, "白龍AIガイド", size=34, bold=True, color="0F2742")
    p = hcell.add_paragraph()
    set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0)
    add_text(p, "企画・場所・待ち時間をすばやく確認できます", size=15, bold=True, color="2B5C8A")

    doc.add_paragraph()

    lead = doc.add_paragraph()
    set_paragraph(lead, WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.15)
    add_text(lead, "画面の案内に沿って操作してください。分からないときは近くのスタッフへ。", size=12.5, bold=True, color="102033")

    # Three feature cards
    cards = doc.add_table(rows=1, cols=3)
    cards.allow_autofit = False
    widths = [Cm(5.85), Cm(5.85), Cm(5.85)]
    card_data = [
        ("企画を探す", "企画名・内容・場所を確認"),
        ("待ち時間を見る", "空いている企画をチェック"),
        ("予定を確認", "ステージや当日の案内を見る"),
    ]
    for i, cell in enumerate(cards.rows[0].cells):
        cards.columns[i].width = widths[i]
        set_cell_shading(cell, "FFFFFF")
        set_cell_border(cell)
        set_cell_margins(cell, top=180, bottom=180, start=130, end=130)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=4)
        add_text(p, card_data[i][0], size=14, bold=True, color="0F2742")
        p2 = cell.add_paragraph()
        set_paragraph(p2, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.05)
        add_text(p2, card_data[i][1], size=10.5, color="334155")

    doc.add_paragraph()

    add_card(
        doc,
        "おすすめの聞き方",
        "「オカダナルドはどこ？」「飲食企画を教えて」「待ち時間が短い企画は？」「今日の外ステージ予定を教えて」",
        fill="F8FAFC",
    )
    doc.add_paragraph()
    add_card(
        doc,
        "大切なお願い",
        "AIの回答は資料をもとにした案内です。最新の変更、売り切れ、雨天時変更、緊急時はスタッフ・先生の案内を優先してください。",
        fill="FFF8E8",
    )

    doc.add_paragraph()

    # QR placeholder and staff note
    bottom = doc.add_table(rows=1, cols=2)
    bottom.allow_autofit = False
    bottom.columns[0].width = Cm(6.2)
    bottom.columns[1].width = Cm(11.8)
    qr = bottom.cell(0, 0)
    note = bottom.cell(0, 1)
    set_cell_shading(qr, "FFFFFF")
    set_cell_border(qr, color="AAB4C3", size="12")
    set_cell_margins(qr, top=520, bottom=520, start=160, end=160)
    p = qr.paragraphs[0]
    set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text(p, "QRコード", size=16, bold=True, color="0F2742")
    p = qr.add_paragraph()
    set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_text(p, "ここに貼付", size=11, color="667085")

    set_cell_shading(note, "F5F7FB")
    set_cell_border(note, color="D7DCE8")
    set_cell_margins(note, top=220, bottom=220, start=220, end=220)
    p = note.paragraphs[0]
    set_paragraph(p, before=0, after=5)
    add_text(p, "スマホで待ち時間を見る場合", size=14, bold=True, color="0F2742")
    p = note.add_paragraph()
    set_paragraph(p, before=0, after=4, line=1.12)
    add_text(p, "掲示されているQRコードを読み取ってください。待ち時間は目安です。実際の列の状況とずれる場合があります。", size=10.5, color="26384D")
    p = note.add_paragraph()
    set_paragraph(p, before=2, after=0, line=1.1)
    add_text(p, "困ったときは、近くの白龍祭スタッフに声をかけてください。", size=11, bold=True, color="0F2742")

    footer = doc.add_paragraph()
    set_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=0)
    add_text(footer, "白龍AIガイドは案内を手伝うためのシステムです。最終案内は当日のスタッフ・先生の指示を優先してください。", size=8.5, color="667085")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
